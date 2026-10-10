"""Attach S2 orphan Barcode2 (RefCode) onto the already-matched szl Item.

Adapted from attach_s6_orphan_barcode2.py, but plans straight from
s2ItemMasterho.xlsx instead of an audit output file.

An "orphan" is a source row whose ItemCode (Barcode1) resolves to exactly one
existing Item while its RefCode (Barcode2) resolves to NO Item (GTIN
leading-zero variants tried). The RefCode is added as an extra Item Barcode on
that Item, copying the UOM of its existing barcode rows. stock_uom, UOM
Conversion Detail and existing barcode rows are never changed.

HELD (never touched here, handled at the very end): rows whose two barcodes
resolve to two different Items, and Items hit by more than one file row.
Scope is every such row regardless of stock (user: "orphan barcode2 kar do").

    ns={}
    path="/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/attach_s2_orphan_barcode2.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()            # dry-run
    ns["main"](apply=True)  # attach on szl, only after backup + approval
"""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

import frappe
from openpyxl import load_workbook

TARGET_SITE = "szl"
SOURCE = "/home/nabeel/frappe-bench/sites/szl/private/files/s2ItemMasterho.xlsx"


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing: expected {TARGET_SITE}, got {site}")
	if not Path(SOURCE).is_file():
		frappe.throw(f"Source file missing: {SOURCE}")


def _norm(v) -> str:
	if v is None:
		return ""
	s = str(v).strip()
	if s.endswith(".0") and s.replace(".", "", 1).isdigit():
		s = s[:-2]
	return "" if s.upper() == "NULL" else s


def _variants(code: str) -> set[str]:
	out = {code, code.upper(), code.lower()}
	if code.isdigit():
		stripped = code.lstrip("0") or "0"
		out.add(stripped)
		for width in (12, 13, 14):
			out.add(code.zfill(width))
			out.add(stripped.zfill(width))
	return {v for v in out if v}


def _indexes():
	barcode_exact, item_exact, variant = defaultdict(set), defaultdict(set), defaultdict(set)
	for barcode, parent in frappe.db.sql(
		"SELECT barcode, parent FROM `tabItem Barcode` WHERE IFNULL(barcode, '') != ''"
	):
		code = str(barcode).strip()
		if code:
			barcode_exact[code.lower()].add(parent)
			for c in _variants(code):
				variant[c.lower()].add(parent)
	for item_code in frappe.get_all("Item", pluck="name"):
		code = str(item_code).strip()
		item_exact[code.lower()].add(item_code)
		for c in _variants(code):
			variant[c.lower()].add(item_code)
	return barcode_exact, item_exact, variant


def _resolve(code, be, ie, vr) -> set[str]:
	if not code:
		return set()
	if be.get(code.lower()):
		return set(be[code.lower()])
	if ie.get(code.lower()):
		return set(ie[code.lower()])
	out = set()
	for c in _variants(code):
		out.update(vr.get(c.lower(), set()))
	return out


def _barcode_uom_for_new_row(item_code: str) -> str:
	uoms = {
		str(u).strip()
		for u in frappe.get_all("Item Barcode", filters={"parent": item_code}, pluck="uom")
		if u and str(u).strip()
	}
	if len(uoms) > 1:
		return None  # mixed Box/Pcs barcodes: never guess, hold for the user
	if uoms:
		return next(iter(uoms))
	stock_uom = frappe.db.get_value("Item", item_code, "stock_uom")
	if not stock_uom:
		frappe.throw(f"{item_code} has no barcode UOM and no stock_uom.")
	return stock_uom


def plan_rows():
	be, ie, vr = _indexes()
	ws = load_workbook(SOURCE, read_only=True, data_only=True).active
	it = ws.iter_rows(values_only=True)
	header = [str(h).strip() if h is not None else "" for h in next(it)]
	candidates, rows_per_item, stats = [], Counter(), Counter()
	for raw in it:
		if not any(v is not None for v in raw):
			continue
		rec = {header[i]: (raw[i] if i < len(raw) else None) for i in range(len(header))}
		b1, b2 = _norm(rec.get("ItemCode")), _norm(rec.get("RefCode"))
		if b2 == b1:
			b2 = ""
		p1, p2 = _resolve(b1, be, ie, vr), _resolve(b2, be, ie, vr)
		item = next(iter(p1 | p2)) if len(p1 | p2) == 1 else None
		for i in p1 | p2:
			rows_per_item[i] += 1
		if not b2 or p2:
			continue  # no Barcode2, or it already exists somewhere
		stats["orphan_candidates"] += 1
		if len(p1) != 1:
			stats["not_applicable_item_not_on_szl"] += 1
			continue
		candidates.append(
			{
				"item": next(iter(p1)),
				"barcode2": b2,
				"onhand": rec.get("Onhand"),
				"desc": _norm(rec.get("Description")),
			}
		)
	planned = []
	seen_barcodes = Counter(c["barcode2"].lower() for c in candidates)
	for c in candidates:
		item = c["item"]
		if rows_per_item[item] > 1:
			planned.append({**c, "action": "held_item_hit_by_multiple_rows"})
			continue
		if seen_barcodes[c["barcode2"].lower()] > 1:
			planned.append({**c, "action": "held_barcode_repeated_in_file"})
			continue
		on_item = {
			str(r).strip().lower()
			for r in frappe.get_all("Item Barcode", filters={"parent": item}, pluck="barcode")
		}
		if c["barcode2"].lower() in on_item:
			planned.append({**c, "action": "already_on_item"})
			continue
		uom = _barcode_uom_for_new_row(item)
		if not uom:
			planned.append({**c, "action": "held_mixed_barcode_uoms"})
			continue
		planned.append({**c, "action": "attach", "uom": uom})
	return planned, stats


def apply_attach(planned):
	attached = 0
	for rec in planned:
		if rec["action"] != "attach":
			continue
		item_code = rec["item"]
		before_uom = frappe.db.get_value("Item", item_code, "stock_uom")
		before = {
			r.barcode: r.uom
			for r in frappe.get_all("Item Barcode", filters={"parent": item_code}, fields=["barcode", "uom"])
		}
		if frappe.db.exists("Item Barcode", {"barcode": rec["barcode2"]}):
			continue  # idempotent re-run / claimed meanwhile
		doc = frappe.get_doc("Item", item_code)
		doc.append("barcodes", {"barcode": rec["barcode2"], "uom": rec["uom"]})
		doc.save(ignore_permissions=True)
		if frappe.db.get_value("Item", item_code, "stock_uom") != before_uom:
			frappe.throw(f"{item_code} stock_uom changed")
		after = {
			r.barcode: r.uom
			for r in frappe.get_all("Item Barcode", filters={"parent": item_code}, fields=["barcode", "uom"])
		}
		for code, uom in before.items():
			if after.get(code) != uom:
				frappe.throw(f"{item_code} existing barcode {code} UOM changed")
		if after.get(rec["barcode2"]) != rec["uom"]:
			frappe.throw(f"{item_code} new barcode not saved with expected UOM")
		attached += 1
		if attached % 100 == 0:
			frappe.db.commit()
	frappe.db.commit()
	return attached


def main(apply: bool = False):
	_assert_site()
	planned, stats = plan_rows()
	counts = Counter(r["action"] for r in planned)
	print(
		{
			"site": TARGET_SITE,
			"apply": apply,
			"source_stats": dict(stats),
			"action_counts": dict(counts),
			"attach_with_stock": sum(1 for r in planned if r["action"] == "attach" and r["onhand"]),
			"sample_attach": [
				(r["item"], r["barcode2"], r["uom"]) for r in planned if r["action"] == "attach"
			][:5],
			"sample_held": [(r["item"], r["barcode2"], r["action"]) for r in planned if r["action"] != "attach"][:8],
		}
	)
	if not apply:
		print("DRY_RUN -- no barcodes attached.")
		return {"counts": counts}
	attached = apply_attach(planned)
	print("ATTACHED", attached)
	return {"attached": attached, "counts": counts}
