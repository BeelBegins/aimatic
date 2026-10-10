"""Create the S2 (Bhatta Chowk) Items missing from szl. Item master only.

Source: sites/szl/private/files/s2ItemMasterho.xlsx (one file: master + S2 stock).
Adapted from create_s6_missing_items.py. Hard-target: live szl. Dry-run by
default; pass apply=True to insert. No warehouse / item_defaults, no prices, no
stock (those are the next pass). ERPNext item_code comes from the naming series.

User decisions 2026-10-11:
- Create ONLY rows with Onhand != 0 (positive or negative). Zero-stock rows are
  ignored entirely.
- Aliases to existing Item Groups: COLOR COSMATICS -> Colour Cosmetics,
  ELECTONIC ITEMS -> Electronic Items, HOUSEHOLD SUNDRIES -> Household
  Essentials, WALL CLOCKS -> Household Essentials, BAGS -> Ladies Bag.
- Junk SubCatName (blank / NULL / TESTING): skipped.
- MRP: where rp exists, rp wins: custom_mrp = rp + rp * Tax_Rate / 100 (item's own
  rate, not a fixed 1.18). No rp: source MRP as-is (0 stays 0).
- Fbr_Tax_Category is the tax source; Tax_Rate is only used for the rp rule.
- HELD, not touched here: rows whose two barcodes resolve to two different
  existing Items, and existing Items hit by two file rows.
- No new Item Group / FBR Tax Category / Brand is ever created.

Run:
    ns={}
    path="/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/create_s2_missing_items.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()           # dry-run
    ns["main"](apply=True) # insert on szl
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import frappe
from openpyxl import load_workbook

TARGET_SITE = "szl"
ITEM_NAMING_SERIES = "STO-ITEM-.YYYY.-"
STOCK_UOM = "Pcs"
SOURCE = "/home/nabeel/frappe-bench/sites/szl/private/files/s2ItemMasterho.xlsx"
SKIP_SUBCATS = {"", "NULL", "TESTING"}

GROUP_ALIASES = {
	"COLOR COSMATICS": "Colour Cosmetics",
	"ELECTONIC ITEMS": "Electronic Items",
	"HOUSEHOLD SUNDRIES": "Household Essentials",
	"WALL CLOCKS": "Household Essentials",
	"BAGS": "Ladies Bag",
}


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing S2 item create: expected site '{TARGET_SITE}', got '{site}'")
	if not Path(SOURCE).is_file():
		frappe.throw(f"Source file missing: {SOURCE}")


def _norm(v) -> str:
	if v is None:
		return ""
	s = str(v).strip()
	if s.endswith(".0") and s.replace(".", "", 1).isdigit():
		s = s[:-2]
	return "" if s.upper() == "NULL" else s


def _num(v) -> float:
	try:
		return float(str(v or "").strip() or 0)
	except ValueError:
		return 0.0


def _variants(code: str) -> set[str]:
	out = {code}
	if code.isdigit():
		stripped = code.lstrip("0") or "0"
		out.add(stripped)
		for n in (12, 13, 14):
			out.add(code.zfill(n))
			out.add(stripped.zfill(n))
	return out


def _title_case(s: str) -> str:
	s = " ".join(str(s or "").split())
	return "" if not s or s.upper() == "NULL" else s.title()


def _catalog():
	idx = {}
	for r in frappe.get_all("Item Barcode", fields=["barcode", "parent"]):
		if r.barcode:
			for v in _variants(str(r.barcode).strip()):
				idx.setdefault(v.lower(), r.parent)
	for name in frappe.get_all("Item", pluck="name"):
		for v in _variants(name):
			idx.setdefault(v.lower(), name)
	return idx


def _resolve(code, idx):
	for v in _variants(code):
		if v.lower() in idx:
			return idx[v.lower()]
	return None


def _mrp(rec) -> float:
	rp = _num(rec.get("rp"))
	if rp:
		return round(rp + rp * _num(rec.get("Tax_Rate")) / 100, 2)
	return _num(rec.get("MRP"))


def _source_rows():
	ws = load_workbook(SOURCE, read_only=True, data_only=True).active
	it = ws.iter_rows(values_only=True)
	header = [str(h).strip() if h is not None else "" for h in next(it)]
	for row in it:
		if any(v is not None for v in row):
			yield {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}


def _plan():
	idx = _catalog()
	groups = {}
	for g in frappe.get_all("Item Group", fields=["name", "item_group_name"]):
		for key in (g.name, g.item_group_name):
			if key:
				groups.setdefault(str(key).lower(), g.name)
	fbr = {n.lower(): n for n in frappe.get_all("FBR Tax Category", pluck="name")}
	brands = {b.lower(): b for b in frappe.get_all("Brand", pluck="name")}

	ready, blocked, c = [], [], Counter()
	seen_new = set()
	for rec in _source_rows():
		c["rows"] += 1
		b1, b2 = _norm(rec.get("ItemCode")), _norm(rec.get("RefCode"))
		if b2 == b1:
			b2 = ""
		i1 = _resolve(b1, idx) if b1 else None
		i2 = _resolve(b2, idx) if b2 else None
		if i1 or i2:
			c["exists_on_szl"] += 1
			continue
		if not _num(rec.get("Onhand")):
			c["zero_stock_ignored"] += 1
			continue
		sub = _norm(rec.get("SubCatName"))
		if sub.upper() in SKIP_SUBCATS:
			c["skipped_junk_subcat"] += 1
			continue
		group = groups.get(GROUP_ALIASES.get(sub.upper(), sub).lower())
		cat = _norm(rec.get("Fbr_Tax_Category"))
		tax = fbr.get(cat.lower()) if cat else None
		if not group:
			blocked.append({"b1": b1, "reason": f"unknown_subcat:{sub}"})
			continue
		if not tax:
			blocked.append({"b1": b1, "reason": f"unknown_or_blank_fbr:{cat}"})
			continue
		if not b1:
			blocked.append({"b1": b1, "reason": "blank_itemcode"})
			continue
		key = b1
		if key in seen_new:
			blocked.append({"b1": b1, "reason": "duplicate_barcode_within_file"})
			continue
		seen_new.add(key)
		brand = brands.get(_norm(rec.get("BrandName")).lower())
		ready.append(
			{
				"item_name": _title_case(rec.get("Description")) or b1,
				"item_group": group,
				"brand": brand,
				"fbr": tax,
				"mrp": _mrp(rec),
				"barcode1": b1,
				"barcode2": b2,
				"onhand": _num(rec.get("Onhand")),
			}
		)
	return ready, blocked, c, idx


def _create_one(plan, idx) -> str:
	for code in (plan["barcode1"], plan["barcode2"]):
		if code and _resolve(code, idx):  # idempotent re-run
			return _resolve(code, idx)
	item = frappe.new_doc("Item")
	item.naming_series = ITEM_NAMING_SERIES
	item.item_name = plan["item_name"]
	item.item_group = plan["item_group"]
	item.stock_uom = STOCK_UOM
	item.is_stock_item = 1
	item.is_sales_item = 1
	item.is_purchase_item = 1
	item.has_variants = 0
	item.disabled = 0
	if plan["brand"]:
		item.brand = plan["brand"]
	item.custom_fbr_tax_category = plan["fbr"]
	if plan["mrp"]:
		item.custom_mrp = plan["mrp"]
	item.append("barcodes", {"barcode": plan["barcode1"], "uom": STOCK_UOM})
	if plan["barcode2"]:
		item.append("barcodes", {"barcode": plan["barcode2"], "uom": STOCK_UOM})
	item.insert(ignore_permissions=True)
	if item.item_defaults:
		item.set("item_defaults", [])
		item.save(ignore_permissions=True)
	for code in (plan["barcode1"], plan["barcode2"]):
		if code:
			for v in _variants(code):
				idx[v.lower()] = item.name
	return item.name


def main(apply: bool = False):
	_assert_site()
	ready, blocked, c, idx = _plan()
	pos = sum(1 for p in ready if p["onhand"] > 0)
	print(
		{
			"site": TARGET_SITE,
			"apply": apply,
			"counters": dict(c),
			"ready": len(ready),
			"ready_positive": pos,
			"ready_negative": len(ready) - pos,
			"ready_with_second_barcode": sum(1 for p in ready if p["barcode2"]),
			"ready_with_mrp": sum(1 for p in ready if p["mrp"]),
			"blocked": len(blocked),
			"blocked_sample": blocked[:10],
		}
	)
	if blocked:
		print("REFUSING to create: blocked rows remain")
		return {"created": 0, "blocked": blocked, "ready": len(ready)}
	if not apply:
		print("DRY_RUN -- no inserts. Pass apply=True to create on", TARGET_SITE)
		print("sample", ready[:3])
		return {"created": 0, "ready": len(ready), "dry_run": True}

	created, failed = [], []
	for i, plan in enumerate(ready, 1):
		try:
			name = _create_one(plan, idx)
			frappe.db.commit()
			created.append(name)
			if i % 200 == 0:
				print(f"created {i}/{len(ready)} last={name}")
		except Exception as exc:
			frappe.db.rollback()
			failed.append({"barcode": plan["barcode1"], "error": str(exc)})
			print(f"FAILED {plan['barcode1']} -> {exc}")
	wh_set = frappe.db.sql(
		"""select count(*) from `tabItem Default` d inner join tabItem i on i.name=d.parent
		where i.creation >= DATE_SUB(NOW(), INTERVAL 2 HOUR) and ifnull(d.default_warehouse,'') != ''"""
	)[0][0]
	print(
		{
			"created": len(created),
			"failed": len(failed),
			"failed_rows": failed[:20],
			"first": created[:3],
			"last": created[-3:],
			"item_defaults_with_warehouse_on_recent": wh_set,
			"item_count": frappe.db.count("Item"),
		}
	)
	return {"created": len(created), "failed": failed, "names_head": created[:5]}
