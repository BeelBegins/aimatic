"""Attach S6 orphan Barcode2 onto the already-matched Item (live szl).

Adapted from attach_s7_orphan_barcode2.py. Reads the xlsx the audit script
produced directly (S7's version assumed a hand-converted CSV; not needed
here). Does not change Item.stock_uom, UOM Conversion Detail, or existing
barcode rows/UOMs. New row copies UOM from that item's existing barcode rows.

    ns={}
    path="/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/attach_s6_orphan_barcode2.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()
    ns["main"](apply=True)
"""

from __future__ import annotations

from pathlib import Path

import frappe
from openpyxl import load_workbook

TARGET_SITE = "szl"
ORPHAN_XLSX = "/tmp/s6_barcode_audit_20260912_235236/s6_orphan_barcode2_20260912_235236.xlsx"


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing: expected {TARGET_SITE}, got {site}")


def _barcode_uom_for_new_row(item_code: str) -> str:
	uoms = [
		str(uom).strip()
		for uom in frappe.get_all("Item Barcode", filters={"parent": item_code}, pluck="uom")
		if uom and str(uom).strip()
	]
	unique = set(uoms)
	if len(unique) > 1:
		frappe.throw(f"{item_code} has mixed barcode UOMs {sorted(unique)}; refusing to guess.")
	if unique:
		return next(iter(unique))
	stock_uom = frappe.db.get_value("Item", item_code, "stock_uom")
	if not stock_uom:
		frappe.throw(f"{item_code} has no barcode UOM and no stock_uom.")
	return stock_uom



def _latest_orphan_xlsx() -> str:
	matches = sorted(Path("/tmp").glob("s6_barcode_audit_*/s6_orphan_barcode2_*.xlsx"))
	if matches:
		return str(matches[-1])
	if Path(ORPHAN_XLSX).is_file():
		return ORPHAN_XLSX
	frappe.throw("Missing orphan audit output. Run audit_s6_barcodes_vs_szl.py first.")


def load_orphans():
	path = Path(_latest_orphan_xlsx())
	if not path.is_file():
		frappe.throw(f"Missing orphan list {ORPHAN_XLSX}")
	wb = load_workbook(path, read_only=True, data_only=True)
	ws = wb.active
	rows_iter = ws.iter_rows(values_only=True)
	header = [str(h).strip() for h in next(rows_iter)]
	out = []
	for row in rows_iter:
		rec = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}
		out.append(
			{
				"Barcode1": str(rec.get("Barcode1") or "").strip(),
				"missing_Barcode2": str(rec.get("missing_Barcode2") or "").strip(),
				"existing_item": str(rec.get("existing_item") or "").strip(),
				"Description": str(rec.get("Description") or "").strip(),
			}
		)
	wb.close()
	return out


def plan_rows():
	planned = []
	for rec in load_orphans():
		item_code = rec["existing_item"].strip()
		barcode = rec["missing_Barcode2"].strip()
		if not frappe.db.exists("Item", item_code):
			planned.append({**rec, "action": "missing_item"})
			continue
		snapshot = {
			"stock_uom": frappe.db.get_value("Item", item_code, "stock_uom"),
			"barcodes": frappe.get_all(
				"Item Barcode", filters={"parent": item_code}, fields=["barcode", "uom"], order_by="idx"
			),
			"uom_conv": frappe.get_all(
				"UOM Conversion Detail", filters={"parent": item_code}, fields=["uom", "conversion_factor"], order_by="idx"
			),
		}
		on_item = {str(r.barcode).strip() for r in snapshot["barcodes"]}
		if barcode in on_item:
			planned.append({**rec, "action": "already_on_item", "uom": None, "snapshot": snapshot})
			continue
		owner = frappe.db.get_value("Item Barcode", {"barcode": barcode}, "parent")
		if owner and owner != item_code:
			planned.append({**rec, "action": "owned_elsewhere", "owner": owner, "snapshot": snapshot})
			continue
		planned.append(
			{**rec, "action": "attach", "uom": _barcode_uom_for_new_row(item_code), "snapshot": snapshot}
		)
	return planned


def apply_attach(planned):
	attached = 0
	for rec in planned:
		if rec["action"] != "attach":
			continue
		item_code = rec["existing_item"]
		before = rec["snapshot"]
		doc = frappe.get_doc("Item", item_code)
		doc.append("barcodes", {"barcode": rec["missing_Barcode2"], "uom": rec["uom"]})
		doc.save(ignore_permissions=True)
		after_uom = frappe.db.get_value("Item", item_code, "stock_uom")
		if after_uom != before["stock_uom"]:
			frappe.throw(f"{item_code} stock_uom changed {before['stock_uom']!r} -> {after_uom!r}")
		after_rows = frappe.get_all(
			"Item Barcode", filters={"parent": item_code}, fields=["barcode", "uom"], order_by="idx"
		)
		after_uom_by_code = {r.barcode: r.uom for r in after_rows}
		for old in before["barcodes"]:
			if after_uom_by_code.get(old.barcode) != old.uom:
				frappe.throw(f"{item_code} existing barcode {old.barcode} UOM changed")
		if after_uom_by_code.get(rec["missing_Barcode2"]) != rec["uom"]:
			frappe.throw(f"{item_code} new barcode not saved with expected UOM")
		attached += 1
	frappe.db.commit()
	return attached


def main(apply: bool = False):
	_assert_site()
	planned = plan_rows()
	summary = {
		"site": TARGET_SITE,
		"apply": apply,
		"counts": {},
		"rows": [
			{
				"item": r["existing_item"],
				"barcode2": r["missing_Barcode2"],
				"action": r["action"],
				"uom": r.get("uom"),
				"owner": r.get("owner"),
				"desc": r["Description"],
			}
			for r in planned
		],
	}
	for rec in planned:
		summary["counts"][rec["action"]] = summary["counts"].get(rec["action"], 0) + 1
	print(summary)
	if not apply:
		print("DRY_RUN — no barcodes attached.")
		return summary
	blocked = [r for r in planned if r["action"] in ("missing_item", "owned_elsewhere")]
	if blocked:
		frappe.throw(f"Refusing apply with blocked rows: {blocked}")
	summary["attached"] = apply_attach(planned)
	print("ATTACHED", summary["attached"])
	return summary
