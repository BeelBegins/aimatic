"""Create the S4-missing Items on siezal (mock). Master file only.

Adapted from create_s7_missing_items.py. Hard-target: siezal for this first
mock pass (update TARGET_SITE to szl at live cutover, same as the S7 script
was updated in place). DRY_RUN via the apply= kwarg. No warehouse /
item_defaults. No prices, no stock. ERPNext item_code from autoname hook (do
not pre-call make_autoname). Unmatched brands left empty. No new Item Groups
/ FBR / Brands beyond the ones already approved and created below.

User decisions 2026-09-13 (see s4_wallayat_complex.md sign-off log):
- ELECTONIC ITEMS -> Electronic Items, WATCH -> Wrist Watches,
  PET ACCESSORIES -> Animal - Pet Items: approved aliases to existing groups.
- FRUITES & VEGETABLES -> Fruits & Vegetables: approved new leaf Item Group
  under existing parent Food Items (created by
  ensure_s4_fruits_vegetables_group.py, already run on siezal).
- SubCatName TESTING (4 rows) and blank SubCatName (9 rows): excluded, same
  treatment as S7's JUNK_DESCRIPTIONS skip (junk / no group to assign).

Run:
    ns={}
    path="/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/create_s4_missing_items.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()           # dry-run
    ns["main"](apply=True) # insert on siezal
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import frappe
import xlrd
from openpyxl import load_workbook

TARGET_SITE = "siezal"
ITEM_NAMING_SERIES = "STO-ITEM-.YYYY.-"
STOCK_UOM = "Pcs"
MASTER_PATH = "/home/nabeel/frappe-bench/sites/szl/private/files/Ho-MasterItemFiled5b674.xlsx"
S4_XLS = "/home/nabeel/frappe-bench/sites/szl/private/files/1004stockposition.xls"
DRY_RUN = True
SKIP_SUBCATS = {"TESTING", "NULL"}  # junk / blank ("NULL" is the literal text some master rows carry)

EXISTING_GROUP_ALIASES = {
	"HOUSEHOLD SUNDRIES": "Household Essentials",
	"HAJI MOEEN (COSMETICS)": "Colour Cosmetics",
	"TOBACCO": "Tobacco",
	"ELECTONIC ITEMS": "Electronic Items",
	"WATCH": "Wrist Watches",
	"PET ACCESSORIES": "Animal - Pet Items",
	"FRUITES & VEGETABLES": "Fruits & Vegetables",
}


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing S4 item create: expected site '{TARGET_SITE}', got '{site}'")


def _norm_code(v) -> str:
	if v is None:
		return ""
	s = str(v).strip()
	if s.endswith(".0") and s.replace(".", "", 1).isdigit():
		s = s[:-2]
	return s


def _variants(code: str) -> set[str]:
	out = {code}
	if code.isdigit():
		stripped = code.lstrip("0") or "0"
		out.add(stripped)
		for n in (12, 13, 14):
			out.add(code.zfill(n))
			out.add(stripped.zfill(n))
	return out


def _title_case(s: str) -> str:
	s = " ".join(str(s or "").split())
	if not s or s.upper() == "NULL":
		return ""
	return s.title()


def _as_float(v) -> float:
	try:
		return float(str(v or "").strip() or 0)
	except ValueError:
		return 0.0


def _parse_s4_rows() -> list[dict]:
	if not Path(S4_XLS).is_file():
		frappe.throw(f"S4 stock file missing: {S4_XLS}")
	wb = xlrd.open_workbook(S4_XLS)
	sh = wb.sheet_by_index(0)
	rows = []
	for r in range(1, sh.nrows):
		row = sh.row_values(r)
		if not row or not any(str(c).strip() for c in row):
			continue
		while len(row) < 7:
			row.append("")

		def _s(v):
			if isinstance(v, float) and v == int(v):
				return str(int(v))
			return str(v).strip()

		b1, b2 = _norm_code(_s(row[0])), _norm_code(_s(row[5]))
		if not b1 and not b2:
			continue
		rows.append({"Barcode1": b1, "Barcode2": b2, "Description": _s(row[1])})
	return rows


def _load_master_index():
	wb = load_workbook(MASTER_PATH, read_only=True, data_only=True)
	ws = wb.active
	it = ws.iter_rows(values_only=True)
	header = [str(h).strip() if h is not None else "" for h in next(it)]
	index = defaultdict(list)
	for row in it:
		rec = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}
		item = _norm_code(rec.get("ItemCode"))
		ref = _norm_code(rec.get("RefCode"))
		if not item and not ref:
			continue
		rec["_ItemCode"] = item
		rec["_RefCode"] = ref
		for code in (item, ref):
			if code:
				for v in _variants(code):
					index[v].append(rec)
	wb.close()
	return index


def _barcode_catalog():
	existing = {}
	for r in frappe.get_all("Item Barcode", fields=["barcode", "parent"]):
		if r.barcode:
			existing[str(r.barcode).strip()] = r.parent
			existing[str(r.barcode).strip().upper()] = r.parent
	for name in frappe.get_all("Item", pluck="name"):
		for v in _variants(name):
			existing.setdefault(v, name)
			existing.setdefault(v.upper(), name)
	variant_to_item = {}
	for code, parent in list(existing.items()):
		for v in _variants(code):
			variant_to_item.setdefault(v, parent)
			variant_to_item.setdefault(v.upper(), parent)
	return existing, variant_to_item


def _exists(code, variant_to_item) -> bool:
	if not code:
		return False
	return any(v in variant_to_item or v.upper() in variant_to_item for v in _variants(code))


def _resolve_group(subcat: str, group_by_lower: dict[str, str]):
	name = (subcat or "").strip()
	if not name or name.upper() == "NULL":
		return None
	alias = EXISTING_GROUP_ALIASES.get(name.upper())
	if alias:
		return group_by_lower.get(alias.lower())
	return group_by_lower.get(name.lower())


def _resolve_fbr(cat: str, fbr_set: set[str], fbr_lower: dict[str, str]):
	name = (cat or "").strip()
	if not name or name.upper() == "NULL":
		return None
	if name in fbr_set:
		return name
	return fbr_lower.get(name.lower())


def _resolve_brand(brand: str, brand_exact: dict[str, str]):
	name = (brand or "").strip()
	if not name or name.upper() == "NULL":
		return None
	if name in brand_exact:
		return brand_exact[name]
	low = name.lower()
	for k, v in brand_exact.items():
		if k.lower() == low:
			return v
	return None


def _resolve_mrp(rec) -> float:
	mrp = _as_float(rec.get("MRP"))
	if mrp:
		return mrp
	rp = _as_float(rec.get("rp"))
	if rp:
		return round(rp * 1.18, 2)
	return 0.0


def _plan_rows():
	s4 = _parse_s4_rows()
	master_index = _load_master_index()
	_existing, variant_to_item = _barcode_catalog()

	groups = frappe.get_all("Item Group", fields=["name", "item_group_name"])
	group_by_lower = {}
	for g in groups:
		for key in (g.name, g.item_group_name):
			if key:
				group_by_lower.setdefault(str(key).lower(), g.name)
	fbr_names = frappe.get_all("FBR Tax Category", pluck="name")
	fbr_set = set(fbr_names)
	fbr_lower = {n.lower(): n for n in fbr_names}
	brand_exact = {b: b for b in frappe.get_all("Brand", pluck="name")}

	ready, skipped, blocked = [], [], []
	for m in s4:
		b1, b2 = m["Barcode1"], m["Barcode2"]
		if _exists(b1, variant_to_item) or _exists(b2, variant_to_item):
			continue
		codes = [c for c in (b1, b2) if c]
		hits, seen = [], set()
		for c in codes:
			for v in _variants(c):
				for rec in master_index.get(v, []):
					if id(rec) not in seen:
						seen.add(id(rec))
						hits.append(rec)
		if not hits:
			blocked.append({**m, "reason": "not_in_master"})
			continue
		exact = [h for h in hits if h["_ItemCode"] in codes or h["_RefCode"] in codes]
		hits = exact or hits
		if len(hits) > 1:
			blocked.append({**m, "reason": f"ambiguous_master:{len(hits)}"})
			continue
		rec = hits[0]
		subcat_raw = str(rec.get("SubCatName") or "").strip()
		if not subcat_raw or subcat_raw.upper() in SKIP_SUBCATS:
			skipped.append({**m, "reason": f"skip_subcat:{subcat_raw or 'blank'}"})
			continue
		group = _resolve_group(subcat_raw, group_by_lower)
		fbr = _resolve_fbr(str(rec.get("Fbr_Tax_Category") or ""), fbr_set, fbr_lower)
		if not group:
			blocked.append({**m, "reason": f"unknown_subcat:{subcat_raw}"})
			continue
		if not fbr:
			blocked.append({**m, "reason": f"unknown_or_blank_fbr:{rec.get('Fbr_Tax_Category')}"})
			continue
		first = rec["_ItemCode"] or b1
		second = rec["_RefCode"]
		if second == first:
			second = ""
		ready.append(
			{
				"s4_b1": b1,
				"s4_b2": b2,
				"item_name": _title_case(str(rec.get("Description") or m["Description"])),
				"item_group": group,
				"brand": _resolve_brand(str(rec.get("BrandName") or ""), brand_exact),
				"fbr": fbr,
				"mrp": _resolve_mrp(rec),
				"barcode1": first,
				"barcode2": second,
			}
		)
	return ready, skipped, blocked, variant_to_item


def _create_one(plan, variant_to_item) -> str:
	for code in (plan["barcode1"], plan["barcode2"]):
		if code and _exists(code, variant_to_item):
			return variant_to_item[next(v for v in _variants(code) if v in variant_to_item)]

	item = frappe.new_doc("Item")
	item.naming_series = ITEM_NAMING_SERIES
	item.item_name = plan["item_name"] or plan["barcode1"]
	item.item_group = plan["item_group"]
	item.stock_uom = STOCK_UOM
	item.is_stock_item = 1
	item.is_sales_item = 1
	item.is_purchase_item = 1
	item.has_variants = 0
	item.disabled = 0
	if plan["brand"]:
		item.brand = plan["brand"]
	item.custom_fbr_tax_category = plan["fbr"]
	if plan["mrp"]:
		item.custom_mrp = plan["mrp"]
	if plan["barcode1"]:
		item.append("barcodes", {"barcode": plan["barcode1"], "uom": STOCK_UOM})
	if plan["barcode2"] and plan["barcode2"] != plan["barcode1"]:
		item.append("barcodes", {"barcode": plan["barcode2"], "uom": STOCK_UOM})
	# Never set item_defaults / warehouse.
	item.insert(ignore_permissions=True)
	if item.item_defaults:
		item.set("item_defaults", [])
		item.save(ignore_permissions=True)
	for code in (plan["barcode1"], plan["barcode2"]):
		if code:
			for v in _variants(code):
				variant_to_item[v] = item.name
	return item.name


def main(apply: bool = False):
	_assert_site()
	if not Path(MASTER_PATH).is_file():
		frappe.throw(f"Master file missing: {MASTER_PATH}")
	ready, skipped, blocked, variant_to_item = _plan_rows()
	print(
		{
			"site": TARGET_SITE,
			"apply": apply,
			"ready": len(ready),
			"skipped": len(skipped),
			"blocked": len(blocked),
			"blocked_sample": blocked[:10],
			"skipped_sample": skipped[:15],
		}
	)
	if blocked:
		print("REFUSING to create: blocked rows remain")
		return {"created": 0, "failed": blocked, "ready": len(ready)}
	if not apply:
		print("DRY_RUN — no inserts. Pass apply=True to create on", TARGET_SITE)
		print("sample", ready[:3])
		return {"created": 0, "ready": len(ready), "dry_run": True}

	created, failed = [], []
	for i, plan in enumerate(ready, 1):
		try:
			name = _create_one(plan, variant_to_item)
			frappe.db.commit()
			created.append(name)
			if i % 200 == 0:
				print(f"created {i}/{len(ready)} last={name}")
		except Exception as exc:
			frappe.db.rollback()
			failed.append({"barcode": plan["barcode1"], "error": str(exc)})
			print(f"FAILED {plan['barcode1']} -> {exc}")
	wh_set = frappe.db.sql(
		"""
		select count(*) from `tabItem Default` d
		inner join tabItem i on i.name=d.parent
		where i.creation >= DATE_SUB(NOW(), INTERVAL 1 HOUR)
		  and ifnull(d.default_warehouse,'') != ''
		"""
	)[0][0]
	series = frappe.db.sql("select name, current from tabSeries where name like %s", ("STO-ITEM%",))
	print(
		{
			"created": len(created),
			"failed": len(failed),
			"failed_rows": failed[:20],
			"first": created[:3],
			"last": created[-3:],
			"item_defaults_with_warehouse_on_recent": wh_set,
			"series": series,
			"item_count": frappe.db.count("Item"),
		}
	)
	return {"created": len(created), "failed": failed, "names_head": created[:5]}
