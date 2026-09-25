"""S6 Foodpanda prices from the portal export productsS6.xlsx.

Barcode-matched rows that already exist in ERPNext, **active or inactive**.
Unmatched / ambiguous / zero-price / disabled / non-sales rows are reported,
not priced. If an Item has both active and inactive portal rows, the active
price wins. Conflicting prices among the chosen source are skipped.
Writes only ``S6 - Khalid Block Foodpanda Price List``. Hard-targets
``szl``. PKR only. SHA256-locked source. Dry-run by default.

    ns = {}
    path = "/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/import_s6_foodpanda_prices.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()            # dry-run
    ns["main"](apply=True)  # only after backup + dry-run
"""

from __future__ import annotations

import hashlib
from collections import Counter, defaultdict
from pathlib import Path

import frappe
from openpyxl import load_workbook

from aimatic.price_export.api import _apply_foodpanda_price_updates

TARGET_SITE = "szl"
BRANCH = "S6 - Khalid Block"
PRICE_LIST = "S6 - Khalid Block Foodpanda Price List"
REQUIRED_CURRENCY = "PKR"
SOURCE_XLSX = "/home/nabeel/frappe-bench/sites/szl/private/files/productsS6.xlsx"
SOURCE_FILE_URL = "/private/files/productsS6.xlsx"
SOURCE_SHA256 = "e351494eb3ca3871d599710a49201057184cf47b1e183f3c0b4583c203aeed23"
EXPECTED_SOURCE_ROWS = 15000
PROTECTED_PRICE_LISTS = {
	"S1 - Ghouri Town VIP Foodpanda Price List",
	"S4 - Wallayat Complex Foodpanda Price List",
	"S5 - Sector C Foodpanda Price List",
	"S7 - Empire Heights Foodpanda Price List",
}


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing S6 Foodpanda import: expected {TARGET_SITE}, got {site}")
	path = Path(SOURCE_XLSX)
	if not path.is_file():
		frappe.throw(f"Missing Foodpanda source {SOURCE_XLSX}")
	digest = hashlib.sha256(path.read_bytes()).hexdigest()
	if digest != SOURCE_SHA256:
		frappe.throw(f"S6 Foodpanda source hash changed: expected {SOURCE_SHA256}, got {digest}")
	if not frappe.db.exists("Price List", PRICE_LIST):
		frappe.throw(f"Missing {PRICE_LIST}")
	currency = frappe.db.get_value("Price List", PRICE_LIST, "currency")
	if currency != REQUIRED_CURRENCY:
		frappe.throw(f"Refusing non-PKR list {PRICE_LIST}: {currency}")
	linked = frappe.db.get_value("Branch", BRANCH, "default_foodpanda_price_list")
	if linked != PRICE_LIST:
		frappe.throw(f"Branch {BRANCH} Foodpanda list is {linked!r}, expected {PRICE_LIST}")


def _norm_code(value) -> str:
	if value is None:
		return ""
	text = str(value).strip()
	if text.endswith(".0") and text.replace(".", "", 1).isdigit():
		text = text[:-2]
	return text


def _variants(code: str) -> set[str]:
	out = {code, code.upper(), code.lower()}
	if code.isdigit():
		stripped = code.lstrip("0") or "0"
		out.add(stripped)
		for width in (12, 13, 14):
			out.add(code.zfill(width))
			out.add(stripped.zfill(width))
	return {value for value in out if value}


def _is_active(value) -> bool:
	if value is True or value == 1:
		return True
	return str(value or "").strip().lower() in {"true", "yes", "y", "1", "active"}


def _as_float(value) -> float:
	try:
		return float(str(value or "").strip().replace(",", "") or 0)
	except (TypeError, ValueError):
		return 0.0


def barcode_indexes():
	barcode_exact = defaultdict(set)
	variant = defaultdict(set)
	for barcode, parent in frappe.db.sql(
		"SELECT barcode, parent FROM `tabItem Barcode` WHERE IFNULL(barcode, '') != ''"
	):
		code = str(barcode).strip()
		if not code:
			continue
		barcode_exact[code.lower()].add(parent)
		for candidate in _variants(code):
			variant[candidate.lower()].add(parent)
	for name in frappe.db.sql_list("SELECT name FROM `tabItem` WHERE IFNULL(name, '') != ''"):
		code = str(name).strip()
		if not code or code.upper().startswith("STO-ITEM"):
			continue
		barcode_exact[code.lower()].add(name)
		for candidate in _variants(code):
			variant[candidate.lower()].add(name)
	return barcode_exact, variant


def _resolve_code(code: str, barcode_exact, variant) -> set[str]:
	if not code:
		return set()
	direct = barcode_exact.get(code.lower(), set())
	if direct:
		return set(direct)
	matches = set()
	for candidate in _variants(code):
		matches.update(variant.get(candidate.lower(), set()))
	return matches


def parse_rows():
	wb = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
	ws = wb["Products"]
	it = ws.iter_rows(values_only=True)
	header = [str(h).strip() if h is not None else "" for h in next(it)]
	rows = []
	for row_number, values in enumerate(it, start=2):
		rec = {header[i]: (values[i] if i < len(values) else None) for i in range(len(header))}
		barcodes = [_norm_code(rec.get(f"barcode {i}")) for i in range(1, 15)]
		barcodes = [code for code in barcodes if code]
		rows.append(
			{
				"row": row_number,
				"sku": str(rec.get("sku") or "").strip(),
				"name": str(rec.get("name") or "").strip(),
				"price": round(_as_float(rec.get("price")), 2),
				"active": _is_active(rec.get("active")),
				"barcodes": barcodes,
			}
		)
	wb.close()
	if len(rows) != EXPECTED_SOURCE_ROWS:
		frappe.throw(f"S6 Foodpanda row count changed: expected {EXPECTED_SOURCE_ROWS}, got {len(rows)}")
	return rows


def _resolve_row(row, barcode_exact, variant):
	primary = _resolve_code(row["barcodes"][0], barcode_exact, variant) if row["barcodes"] else set()
	all_matches = set(primary)
	for code in row["barcodes"][1:]:
		all_matches |= _resolve_code(code, barcode_exact, variant)
	if len(primary) == 1:
		return next(iter(primary)), "barcode1", None
	if len(primary) > 1:
		return None, "ambiguous_barcode1", sorted(primary)
	if len(all_matches) == 1:
		return next(iter(all_matches)), "other_barcode", None
	if len(all_matches) > 1:
		return None, "ambiguous_other_barcodes", sorted(all_matches)
	return None, "unmatched", None


def build_plan(rows):
	barcode_exact, variant = barcode_indexes()
	disabled = set(frappe.get_all("Item", filters={"disabled": 1}, pluck="name"))
	not_sales = set(frappe.get_all("Item", filters={"is_sales_item": 0}, pluck="name"))
	stats = Counter()
	issues = []
	unmatched_sample = []
	by_item = {}

	for row in rows:
		stats["source_rows"] += 1
		if row["active"]:
			stats["active"] += 1
		else:
			stats["inactive"] += 1
		if row["price"] <= 0:
			stats["zero_price"] += 1
			continue
		item_code, resolution, extra = _resolve_row(row, barcode_exact, variant)
		if item_code is None:
			stats[resolution] += 1
			if resolution != "unmatched":
				issues.append(
					{
						"row": row["row"],
						"sku": row["sku"],
						"name": row["name"],
						"reason": resolution,
						"items": extra,
						"active": row["active"],
					}
				)
			elif len(unmatched_sample) < 15:
				unmatched_sample.append(
					{
						"row": row["row"],
						"sku": row["sku"],
						"name": row["name"],
						"barcode1": (row["barcodes"] or [""])[0],
						"active": row["active"],
					}
				)
			continue
		stats[resolution] += 1
		if item_code in disabled:
			stats["disabled"] += 1
			issues.append({"row": row["row"], "sku": row["sku"], "reason": "disabled", "items": [item_code]})
			continue
		if item_code in not_sales:
			stats["not_sales_item"] += 1
			issues.append({"row": row["row"], "sku": row["sku"], "reason": "not_sales_item", "items": [item_code]})
			continue
		rec = by_item.setdefault(
			item_code,
			{
				"item_code": item_code,
				"active_prices": set(),
				"inactive_prices": set(),
				"rows": [],
				"skus": [],
			},
		)
		if row["active"]:
			rec["active_prices"].add(row["price"])
			stats["matched_active"] += 1
		else:
			rec["inactive_prices"].add(row["price"])
			stats["matched_inactive"] += 1
		rec["rows"].append(row["row"])
		rec["skus"].append(row["sku"])

	prices = []
	for item_code, rec in by_item.items():
		if rec["active_prices"]:
			chosen = rec["active_prices"]
			source = "active"
		else:
			chosen = rec["inactive_prices"]
			source = "inactive"
		if len(chosen) > 1:
			stats["conflicting_prices"] += 1
			issues.append(
				{
					"reason": "conflicting_prices",
					"item_code": item_code,
					"prices": sorted(chosen),
					"rows": rec["rows"],
					"skus": rec["skus"],
					"source": source,
				}
			)
			continue
		price = next(iter(chosen))
		prices.append({"item_code": item_code, "price": price, "source": source})
		if source == "inactive":
			stats["ready_inactive_only"] += 1
		else:
			stats["ready_active"] += 1
		if len(rec["rows"]) > 1:
			stats["merged_duplicate_item_rows"] += len(rec["rows"]) - 1

	stats["ready_prices"] = len(prices)
	stats["blocking_issues"] = len(issues)
	return {
		"stats": stats,
		"issues": issues,
		"prices": prices,
		"unmatched_sample": unmatched_sample,
	}


def _other_list_counts():
	return {name: frappe.db.count("Item Price", {"price_list": name}) for name in sorted(PROTECTED_PRICE_LISTS)}


def main(apply: bool = False):
	_assert_site()
	rows = parse_rows()
	plan = build_plan(rows)
	print(
		{
			"site": TARGET_SITE,
			"apply": apply,
			"branch": BRANCH,
			"price_list": PRICE_LIST,
			"source": SOURCE_XLSX,
			"source_sha256": SOURCE_SHA256,
			"stats": dict(plan["stats"]),
			"issue_counts": dict(Counter(issue["reason"] for issue in plan["issues"])),
			"issue_sample": plan["issues"][:15],
			"unmatched_sample": plan["unmatched_sample"],
			"current_s6_foodpanda_prices": frappe.db.count("Item Price", {"price_list": PRICE_LIST}),
			"protected_counts": _other_list_counts(),
		}
	)
	if not apply:
		print("DRY_RUN — no Foodpanda prices posted. Ready rows can apply; issues are skipped.")
		return {"applied": False, "blocked": bool(plan["issues"]), "stats": plan["stats"], "ready": len(plan["prices"])}

	before_protected = _other_list_counts()
	normalized = {row["item_code"]: {"price": row["price"]} for row in plan["prices"]}
	keys = list(normalized)
	created = updated = unchanged = skipped_invalid = 0
	logs = []
	chunk_size = 400
	for offset in range(0, len(keys), chunk_size):
		chunk = {code: normalized[code] for code in keys[offset : offset + chunk_size]}
		result = _apply_foodpanda_price_updates(
			BRANCH,
			chunk,
			source_file=SOURCE_FILE_URL,
			skip_invalid=True,
		)
		if result["price_list"] != PRICE_LIST:
			frappe.throw(f"Refusing write: expected {PRICE_LIST}, got {result['price_list']}")
		created += result["created"]
		updated += result["updated"]
		unchanged += result["unchanged"]
		skipped_invalid += result.get("skipped_invalid") or 0
		logs.append(result["log"])
		frappe.db.commit()
		print("CHUNK", offset, result)
	after_protected = _other_list_counts()
	if after_protected != before_protected:
		frappe.throw(f"Protected Foodpanda lists changed: {before_protected} -> {after_protected}")
	summary = {
		"applied": True,
		"price_list": PRICE_LIST,
		"created": created,
		"updated": updated,
		"unchanged": unchanged,
		"skipped_invalid": skipped_invalid,
		"logs": logs,
		"ready": len(normalized),
		"current_s6_foodpanda_prices": frappe.db.count("Item Price", {"price_list": PRICE_LIST}),
		"protected_counts": after_protected,
	}
	print(summary)
	return summary
