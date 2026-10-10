"""S2 (Bhatta Chowk) customers + opening loyalty points. Loyalty ONLY.

Source: sites/szl/private/files/S2CustomerData.xlsx. Adapted from
import_szl_customers.py (see customerimport.md) with the S2 decisions of
2026-10-11:

- No customer balance: CurrentBalance / OpeningBalance / CreditLimit are ignored.
  Only CustomerCode, CustomerName, MobileNo and LoyaltyPoints are used.
- Rows with a mobile group by normalized mobile (the Customer PK); points are
  summed across the group and every code is kept. Rows with NO mobile (the file
  stores the literal text NULL) are each their own Customer, identified by
  CustomerCode.
- A mobile that already belongs to an existing szl Customer reuses that Customer
  and ADDS the S2 points as a new Loyalty Point Entry (not replace).
- New Customers get default_price_list = S2 - Bhatta Chowk Selling Price List
  explicitly. Selling Settings is never touched (it is site-wide, live S1).
- S2 CustomerCodes are stored as "S2:<code>" in custom_legacy_customer_code so
  they can never be confused with S1 codes, and S2 loyalty entries carry a
  distinct reason tag so a customer's existing S1 opening entry does not make
  the S2 one look "already done".
- Run LAST in the S2 migration.

    ns={}
    path="/home/nabeel/frappe-bench/apps/aimatic/ipos_data_migration/import_s2_customers.py"
    exec(compile(open(path).read(), path, "exec"), ns)
    ns["main"]()            # dry-run
    ns["main"](apply=True)  # live, only after backup + approval
"""

from __future__ import annotations

import collections
from pathlib import Path

import frappe
from frappe.utils import add_days, cint, flt, today
from openpyxl import load_workbook

from aimatic.offline_pos.customer_validation import normalize_pak_mobile

TARGET_SITE = "szl"
FILE_PATH = "/home/nabeel/frappe-bench/sites/szl/private/files/S2CustomerData.xlsx"
CODE_PREFIX = "S2:"
REASON_TAG = "Legacy iPOS S2 opening loyalty balance (S2CustomerData import"
CUSTOMER_GROUP = "Individual"
TERRITORY = "Pakistan"
LOYALTY_PROGRAM_NAME = "Siezal Loyalty Program"
LOYALTY_EXPIRY_DAYS = 365
S2_PRICE_LIST = "S2 - Bhatta Chowk Selling Price List"


def _assert_site():
	site = getattr(frappe.local, "site", None) or ""
	if site != TARGET_SITE:
		frappe.throw(f"Refusing: expected site '{TARGET_SITE}', got '{site}'")
	if not Path(FILE_PATH).is_file():
		frappe.throw(f"Source file missing: {FILE_PATH}")
	pl = frappe.db.get_value("Price List", S2_PRICE_LIST, ["enabled", "selling"], as_dict=True)
	if not pl or not pl.enabled or not pl.selling:
		frappe.throw(f"Missing or disabled Selling Price List {S2_PRICE_LIST}")
	if not frappe.db.exists("Loyalty Program", LOYALTY_PROGRAM_NAME):
		frappe.throw(f"Loyalty Program {LOYALTY_PROGRAM_NAME} is missing")
	if not frappe.db.exists("Custom Field", "Customer-custom_legacy_customer_code"):
		frappe.throw("Customer.custom_legacy_customer_code field is missing")


def _clean(v) -> str:
	text = str(v).strip() if v is not None else ""
	return "" if text.upper() == "NULL" else text


def load_rows() -> list[dict]:
	ws = load_workbook(FILE_PATH, read_only=True, data_only=True).active
	it = ws.iter_rows(values_only=True)
	header = [str(h).strip() if h is not None else "" for h in next(it)]
	rows = []
	for n, raw in enumerate(it, start=2):
		if not any(v is not None for v in raw):
			continue
		rec = {header[i]: (raw[i] if i < len(raw) else None) for i in range(len(header))}
		mobile_raw = _clean(rec.get("MobileNo"))
		normalized = normalize_pak_mobile(mobile_raw) if mobile_raw else None
		rows.append(
			{
				"row": n,
				"code": _clean(rec.get("CustomerCode")),
				"name": _clean(rec.get("CustomerName")),
				"mobile_raw": mobile_raw,
				"mobile": normalized,
				"points": flt(rec.get("LoyaltyPoints") or 0),
				"disabled": cint(rec.get("Disabled") or 0),
			}
		)
	return rows


def build_groups(rows):
	"""mobile -> members; null-mobile rows get their own key by code."""
	groups = collections.OrderedDict()
	for r in rows:
		key = r["mobile"] if r["mobile"] else f"code:{r['code']}"
		groups.setdefault(key, []).append(r)
	return groups


def _winner(members):
	return max(members, key=lambda r: abs(r["points"]))


def _legacy(members) -> str:
	return ",".join(f"{CODE_PREFIX}{m['code']}" for m in members)


def _existing_customer(key, members):
	if not key.startswith("code:"):
		found = frappe.db.get_value("Customer", {"mobile_no": key})
		if found:
			return found, "existing_mobile"
	for m in members:
		hit = frappe.db.get_value("Customer", {"custom_legacy_customer_code": ["like", f"%{CODE_PREFIX}{m['code']}%"]})
		if hit:
			# exact-token check so S2:1 never matches S2:10
			codes = (frappe.db.get_value("Customer", hit, "custom_legacy_customer_code") or "").split(",")
			if f"{CODE_PREFIX}{m['code']}" in [c.strip() for c in codes]:
				return hit, "already_imported"
	return None, None


def plan(groups):
	stats = collections.Counter()
	existing_mobile, reused = [], []
	name_collisions = 0
	for key, members in groups.items():
		stats["groups"] += 1
		if key.startswith("code:"):
			stats["no_mobile_groups"] += 1
		if len(members) > 1:
			stats["merged_by_mobile"] += 1
		if any(m["disabled"] for m in members):
			stats["disabled_in_file"] += 1
		pts = sum(m["points"] for m in members)
		if cint(round(pts)):
			stats["entries_planned"] += 1
			if pts < 0:
				stats["negative_point_groups"] += 1
		cust, why = _existing_customer(key, members)
		if cust:
			stats[why] += 1
			(existing_mobile if why == "existing_mobile" else reused).append((key, cust, round(pts, 2)))
		else:
			stats["customers_to_create"] += 1
			w = _winner(members)
			nm = w["name"].title() if w["name"] else (key if not key.startswith("code:") else "")
			if nm and frappe.db.exists("Customer", nm):
				name_collisions += 1
	stats["name_collisions_will_be_suffixed"] = name_collisions
	return stats, existing_mobile, reused


def create_customers(groups):
	key_to_customer, stats, failures = {}, collections.Counter(), []
	for key, members in groups.items():
		w = _winner(members)
		try:
			cust, why = _existing_customer(key, members)
			if cust:
				key_to_customer[key] = cust
				stats["customers_reused"] += 1
				if why == "existing_mobile":
					# keep the S2 code visible on the existing record
					cur = frappe.db.get_value("Customer", cust, "custom_legacy_customer_code") or ""
					tokens = [c.strip() for c in cur.split(",") if c.strip()]
					add = [f"{CODE_PREFIX}{m['code']}" for m in members if f"{CODE_PREFIX}{m['code']}" not in tokens]
					if add:
						frappe.db.set_value("Customer", cust, "custom_legacy_customer_code", ",".join(tokens + add))
						stats["legacy_code_appended_to_existing"] += 1
				frappe.db.commit()
				continue
			name = w["name"].title() if w["name"] else ""
			suffix = key if not key.startswith("code:") else f"{CODE_PREFIX}{w['code']}"
			if not name:
				name = suffix
			elif frappe.db.exists("Customer", name):
				name = f"{name} ({suffix})"
			c = frappe.new_doc("Customer")
			c.customer_name = name
			c.customer_type = "Individual"
			c.customer_group = CUSTOMER_GROUP
			c.territory = TERRITORY
			if not key.startswith("code:"):
				c.mobile_no = key
			c.custom_legacy_customer_code = _legacy(members)
			c.loyalty_program = LOYALTY_PROGRAM_NAME
			c.default_price_list = S2_PRICE_LIST
			c.insert(ignore_permissions=True)
			key_to_customer[key] = c.name
			stats["customers_created"] += 1
			frappe.db.commit()
		except Exception as exc:
			frappe.db.rollback()
			stats["group_failed"] += 1
			failures.append({"key": key, "codes": _legacy(members), "reason": str(exc)})
			print(f"FAILED customer {key} -> {exc}")
	return key_to_customer, stats, failures


def create_entries(groups, key_to_customer, company, posting_date):
	expiry = add_days(posting_date, LOYALTY_EXPIRY_DAYS)
	stats, failures = collections.Counter(), []
	drift = 0.0
	for key, members in groups.items():
		total = sum(m["points"] for m in members)
		pts = cint(round(total))
		if not pts:
			continue
		customer = key_to_customer.get(key)
		if not customer:
			stats["skipped_no_customer"] += 1
			continue
		reason = f"{REASON_TAG}, codes {_legacy(members)})"
		if frappe.db.exists(
			"Loyalty Point Entry",
			{"customer": customer, "loyalty_program": LOYALTY_PROGRAM_NAME, "discretionary_reason": ["like", f"{REASON_TAG}%"]},
		):
			stats["entries_reused"] += 1
			continue
		try:
			e = frappe.new_doc("Loyalty Point Entry")
			e.loyalty_program = LOYALTY_PROGRAM_NAME
			e.customer = customer
			e.loyalty_points = pts
			e.purchase_amount = 0
			e.posting_date = posting_date
			e.expiry_date = expiry
			e.company = company
			e.invoice_type = "Journal Entry"  # placeholder marker, as S1
			e.discretionary_reason = reason
			e.insert(ignore_permissions=True)
			stats["entries_created"] += 1
			drift += pts - total
			frappe.db.commit()
		except Exception as exc:
			frappe.db.rollback()
			stats["entries_failed"] += 1
			failures.append({"customer": customer, "key": key, "reason": str(exc)})
			print(f"FAILED loyalty entry {customer} -> {exc}")
	stats["rounding_drift"] = round(drift, 2)
	return stats, failures


def main(apply: bool = False):
	_assert_site()
	rows = load_rows()
	groups = build_groups(rows)
	pstats, existing_mobile, reused = plan(groups)
	total_points = round(sum(r["points"] for r in rows), 2)
	print(
		{
			"site": TARGET_SITE,
			"apply": apply,
			"rows": len(rows),
			"groups": len(groups),
			"file_total_points": total_points,
			"price_list_for_new_customers": S2_PRICE_LIST,
			"plan": dict(pstats),
			"existing_mobile_sample": existing_mobile[:10],
			"already_imported_sample": reused[:5],
		}
	)
	if not apply:
		print("DRY_RUN -- no Customer or Loyalty Point Entry created.")
		return {"plan": pstats}

	company = frappe.get_all("Company", pluck="name")
	if len(company) != 1:
		frappe.throw(f"Expected exactly one Company, found {company}")
	key_to_customer, cstats, cfail = create_customers(groups)
	print("PHASE 1 customers", dict(cstats), "failures", len(cfail))
	estats, efail = create_entries(groups, key_to_customer, company[0], today())
	print("PHASE 2 loyalty", dict(estats), "failures", len(efail))
	for f in (cfail + efail)[:30]:
		print("FAILURE", f)
	return {"customers": cstats, "loyalty": estats}
