import unittest

from docx import Document

from aimaticlearning.lms_learning.content_format import (
    is_chapter_heading,
    paragraph_kind,
    parse_notes_blocks,
    render_content_blocks,
    render_docx_blocks,
)


class TestContentFormat(unittest.TestCase):
    def test_docx_heading_and_numbered_normal_heading(self):
        document = Document()
        chapter = document.add_paragraph("Chapter 1: Formation")
        chapter.style = "Title"
        section = document.add_paragraph("1. Offer and acceptance")
        body = document.add_paragraph("An offer must be communicated to the offeree.")
        self.assertTrue(is_chapter_heading(chapter))
        self.assertEqual(paragraph_kind(section)["kind"], "heading")
        html = render_docx_blocks(document.paragraphs, chapter_title="Chapter 1: Formation")
        self.assertIn("<h2>Chapter 1: Formation</h2>", html)
        self.assertIn("<h3>1. Offer and acceptance</h3>", html)
        self.assertIn("<p>An offer must be communicated to the offeree.</p>", html)

    def test_render_groups_ordered_and_unordered_items_and_escapes_text(self):
        blocks = [
            {"kind": "heading", "level": 3, "text": "Key points"},
            {"kind": "list_item", "ordered": True, "level": 0, "text": "First <rule>"},
            {"kind": "list_item", "ordered": True, "level": 0, "text": "Second"},
            {"kind": "list_item", "ordered": False, "level": 0, "text": "Bullet"},
        ]
        html = render_content_blocks(blocks)
        self.assertIn("<ol><li>First &lt;rule&gt;</li><li>Second</li></ol>", html)
        self.assertIn("<ul><li>Bullet</li></ul>", html)

    def test_parse_notes_blocks_keeps_heading_and_list_semantics(self):
        blocks = parse_notes_blocks(
            "<h2>Chapter 1</h2><h3>Offer</h3><p>Meaning</p>"
            "<ul><li>One</li><li>Two &amp; two</li></ul>"
        )
        self.assertEqual(
            blocks,
            [
                {"kind": "heading", "level": 2, "text": "Chapter 1"},
                {"kind": "heading", "level": 3, "text": "Offer"},
                {"kind": "paragraph", "text": "Meaning"},
                {"kind": "list", "ordered": False, "items": ["One", "Two & two"]},
            ],
        )

    def test_docx_table_has_accessible_header_and_body(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Rule"
        table.cell(0, 1).text = "Effect"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        html = render_docx_blocks([table])
        self.assertIn("<thead><tr><th>Rule</th><th>Effect</th></tr></thead>", html)
        self.assertIn("<tbody><tr><td>A</td><td>B</td></tr></tbody>", html)


if __name__ == "__main__":
    unittest.main()
